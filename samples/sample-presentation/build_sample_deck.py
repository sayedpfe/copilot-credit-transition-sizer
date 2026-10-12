# Builds the sample presentation from scratch (python-pptx, neutral branding).
# All numbers come from the synthetic demo dataset in ../sample-data — no real
# tenant data, no proprietary slide templates. Rebuild: py build_sample_deck.py
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PRIMARY = RGBColor(0x06, 0x5A, 0x82)
SECONDARY = RGBColor(0x1C, 0x72, 0x93)
ACCENT = RGBColor(0x21, 0x29, 0x5C)
SURFACE = RGBColor(0xF5, 0xF7, 0xFA)
TEXT = RGBColor(0x21, 0x21, 0x21)
MUTED = RGBColor(0x5A, 0x67, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xF5, 0x7F, 0x17)
RED = RGBColor(0xC6, 0x28, 0x28)
BORD = RGBColor(0xD6, 0xDC, 0xE6)
F = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
PAGE = [0]


def slide():
    s = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = SURFACE
    return s


def tb(s, x, y, w, h, text, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
       font=F, anchor=MSO_ANCHOR.TOP):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold; r.font.name = font
    return sh


def box(s, x, y, w, h, fill=SURFACE, line=BORD, text="", size=13, color=TEXT,
        bold=False, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    if text:
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, ln in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run(); r.text = ln
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = F
    return sh


def title(s, text):
    tb(s, 0.5, 0.32, 12.3, 0.7, text, size=30, color=PRIMARY, bold=True)


def footer(s):
    tb(s, 0.5, 7.08, 10.5, 0.32,
       "copilot-credit-transition-sizer · synthetic demo data · not an official Microsoft solution",
       size=10, color=MUTED)
    tb(s, 12.4, 7.08, 0.45, 0.32, str(PAGE[0]), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


# S1 Title
s = slide()
box(s, 0, 0, 13.333, 7.5, fill=ACCENT, line=None, shape=MSO_SHAPE.RECTANGLE)
tb(s, 0.8, 2.3, 11.7, 1.6, "Copilot Credit Transition Sizer", size=44, color=WHITE, bold=True)
tb(s, 0.8, 3.6, 11.5, 0.5, "Size the Copilot Credit purchase from actual usage — before seeded AI Builder credits retire",
   size=18, color=RGBColor(0xC9, 0xD4, 0xE0))
box(s, 0.8, 4.6, 9.6, 0.6, fill=AMBER, line=None,
    text="Seeded AI Builder credits end Nov 1, 2026 (or contract term) — with no automatic conversion",
    size=14, color=WHITE, bold=True)
tb(s, 0.8, 6.9, 11, 0.4, "Open source · MIT · all numbers in this deck are synthetic demo data", size=11,
   color=RGBColor(0x9F, 0xAC, 0xC4))

# S2 No flat conversion
s = slide()
title(s, "There is no flat conversion — the capability decides the price")
tb(s, 0.6, 1.1, 12, 0.4, "What the same 1,000,000 AI Builder credits becomes in Copilot Credits:", size=14, color=MUTED)
bars = [("Object detection (the only 1:1)", 1000000, "1,000,000", PRIMARY),
        ("Invoice / receipt / ID processing", 250000, "250,000", SECONDARY),
        ("Prompts — basic (blended)", 83333, "≈ 83,300", SECONDARY),
        ("Custom document processing", 80000, "80,000", SECONDARY),
        ("Text recognition (OCR)", 33333, "≈ 33,300", SECONDARY)]
y = 1.75
for label, v, vt, c in bars:
    tb(s, 0.6, y + 0.03, 3.4, 0.5, label, size=12, align=PP_ALIGN.RIGHT)
    w = max(0.07, 6.6 * v / 1000000)
    box(s, 4.15, y, w, 0.38, fill=c, line=None, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 4.15 + w + 0.12, y + 0.01, 2.0, 0.4, vt, size=13, color=PRIMARY, bold=True)
    y += 0.6
panel = box(s, 0.6, 4.95, 12.1, 1.8, fill=WHITE, line=BORD)
tf = panel.text_frame; tf.word_wrap = True
tf.margin_left = Pt(10); tf.margin_top = Pt(8)
for i, t in enumerate(["Microsoft publishes no conversion ratio between the two currencies",
                       "Per capability: factor = Copilot Credit rate ÷ AI Builder rate (×0.033 … ×1.0)",
                       "A flat ratio would mis-size badly — the tool re-rates every single event instead",
                       "Rates live in versioned JSON config with source URLs — a Microsoft change is a config edit"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = "•  " + t
    r.font.size = Pt(13); r.font.name = F; r.font.color.rgb = TEXT
    p.space_after = Pt(5)
footer(s)

# S3 How it works
s = slide()
title(s, "How the tool works — measured, not estimated")
box(s, 0.6, 2.4, 2.8, 1.5, fill=WHITE, line=PRIMARY, text="Dataverse\nAI Event table\n(every AI run)", size=13)
box(s, 3.75, 2.55, 0.7, 1.2, fill=None, line=None, text="→", size=28, color=PRIMARY, bold=True)
box(s, 4.45, 2.3, 3.3, 1.7, fill=PRIMARY, line=None,
    text="sizer (read-only)\ntags currency per event\nre-rates 100% into\nCopilot Credits", size=13, color=WHITE, bold=True)
box(s, 7.85, 2.55, 0.7, 1.2, fill=None, line=None, text="→", size=28, color=PRIMARY, bold=True)
box(s, 8.55, 2.4, 3.3, 1.5, fill=WHITE, line=SECONDARY,
    text="Sizing reports\nper flow · per user · per env\n→ packs to buy", size=13)
tb(s, 4.45, 4.15, 3.3, 0.4, "peak month + growth buffer", size=11, color=MUTED, align=PP_ALIGN.CENTER)
for i, (t, d) in enumerate([
        ("Zero-config auth", "az login as admin — no app registration"),
        ("Read-only", "GET requests only; nothing changes"),
        ("Fallback without API", "PPAC consumption-report CSV mode")]):
    x = 0.6 + i * 4.1
    card = box(s, x, 5.1, 3.9, 1.3, fill=WHITE, line=BORD)
    tf = card.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(10)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = t
    r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = PRIMARY; r.font.name = F
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = d
    r2.font.size = Pt(12); r2.font.color.rgb = TEXT; r2.font.name = F
footer(s)

# S4 Six phases
s = slide()
title(s, "The transition journey")
phases = [("1 MEASURE", "run the sizer", SECONDARY), ("2 VALIDATE", "reconcile vs PPAC ±5%", SECONDARY),
          ("3 PURCHASE", "buy packs (M365 admin)", PRIMARY), ("4 ALLOCATE", "credits per environment", PRIMARY),
          ("5 GUARDRAILS", "limits · alerts · PAYG net", PRIMARY), ("6 CUTOVER", "automatic — no rebuild", GREEN)]
x = 0.45
for name, sub, c in phases:
    box(s, x, 2.9, 2.05, 1.05, fill=c, line=None, text=name, size=13, color=WHITE, bold=True, shape=MSO_SHAPE.CHEVRON)
    tb(s, x + 0.1, 4.15, 2.0, 0.85, sub, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    x += 2.09
tb(s, 0.6, 5.6, 12.1, 0.5, "Apps and flows fall back to Copilot Credits automatically — the cutover is a licensing event, not a migration.",
   size=14, color=TEXT, align=PP_ALIGN.CENTER)
footer(s)

# S5 Sample output
s = slide()
title(s, "Sample output (synthetic demo tenant)")
mono = box(s, 0.6, 1.3, 12.1, 4.3, fill=RGBColor(0x1E, 0x24, 0x30), line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tf = mono.text_frame; tf.word_wrap = False
tf.margin_left = Pt(14); tf.margin_top = Pt(10)
lines = [
    "=== Copilot Credit Transition Sizing ===",
    "  contoso-demo-prod",
    "    peak month  : 2025-12  ->  30,953 CC forecast",
    "    with buffer : 37,143 CC  ->  2 pack(s)",
    "      - Invoice intake flow: ~7,133 CC/month (peak 24,000 in 2025-12)",
    "      - Same-credits demo: object detection: ~3,200 CC/month",
    "  contoso-demo-dev",
    "    peak month  : 2026-03  ->  10,528 CC  ->  1 pack(s)",
    "",
    "  TENANT peak: 33,160 CC -> +20% buffer -> 2 pack(s) of 25,000 CC/month",
    "  Sum of per-environment packs: 3   (pooling saves 1 pack)",
]
for i, ln in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = ln if ln else " "
    r.font.size = Pt(14); r.font.name = "Consolas"
    r.font.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    p.space_after = Pt(3)
tb(s, 0.6, 5.8, 12.1, 0.8,
   "Full sample outputs ship in samples/sample-output: the Excel workbook (Summary · Per-Flow · Per-User · Per-Environment ·\nPer-Capability · Unmapped · How-it-is-calculated) and the executive markdown report.",
   size=12, color=MUTED)
footer(s)

# S6 Trust gate
s = slide()
title(s, "The trust gate — validate before anyone buys")
box(s, 0.9, 2.3, 4.6, 1.5, fill=WHITE, line=PRIMARY, text="sizer result\nAI Builder credits per\nenvironment, current month", size=13)
box(s, 7.8, 2.3, 4.6, 1.5, fill=WHITE, line=SECONDARY, text="PPAC consumption report\n(Microsoft's own numbers,\nadmin download)", size=13)
box(s, 5.65, 2.75, 2.0, 0.6, fill=None, line=None, text="compare", size=14, color=MUTED)
box(s, 2.4, 4.3, 3.7, 0.55, fill=GREEN, line=None, text="delta ≤ 5%  →  proceed", size=13, color=WHITE, bold=True)
box(s, 7.3, 4.3, 3.7, 0.55, fill=AMBER, line=None, text="delta > 5%  →  FLAGGED + guidance", size=13, color=WHITE, bold=True)
tb(s, 0.9, 5.4, 11.6, 0.8,
   "PPAC computation can lag 5+ days (Microsoft-documented) — the tool says so and recommends comparing a closed month.\nThe sample data includes a deliberately lagging PPAC file so you can rehearse the flagged path too.",
   size=12, color=MUTED, align=PP_ALIGN.CENTER)
footer(s)

# S7 Get started + disclaimer
s = slide()
title(s, "Get started")
mono = box(s, 0.6, 1.4, 7.6, 2.9, fill=RGBColor(0x1E, 0x24, 0x30), line=None)
tf = mono.text_frame; tf.margin_left = Pt(14); tf.margin_top = Pt(10)
for i, ln in enumerate(["git clone <repo> && cd copilot-credit-transition-sizer",
                        "npm ci && npm run build && npm test",
                        "node tools/generate-demo-data.mjs          # rehearse",
                        "node dist/cli.js --data-dir ./data-demo \\",
                        "     --out-dir ./output-demo forecast",
                        "az login --allow-no-subscriptions          # real run",
                        "node dist/cli.js ingest --auth azcli \\",
                        "     --environment https://org.crm.dynamics.com"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = ln
    r.font.size = Pt(13); r.font.name = "Consolas"; r.font.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    p.space_after = Pt(3)
card = box(s, 8.5, 1.4, 4.2, 2.9, fill=WHITE, line=BORD)
tf = card.text_frame; tf.word_wrap = True; tf.margin_left = Pt(10); tf.margin_top = Pt(8)
for i, t in enumerate(["RUN-GUIDE.md — validated step-by-step",
                       "docs/solution-architecture.html — full design",
                       "README — Microsoft Learn citations + rate-refresh guide",
                       "samples/ — data, reports, this deck"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = "•  " + t
    r.font.size = Pt(12.5); r.font.name = F; r.font.color.rgb = TEXT
    p.space_after = Pt(6)
box(s, 0.6, 4.8, 12.1, 1.4, fill=SURFACE, line=AMBER,
    text="Disclaimer: estimates only — not an official Microsoft solution. Verify every figure against the\nMicrosoft Power Platform Licensing Guide and the Copilot Credit Guide before purchasing.",
    size=13, color=TEXT)
footer(s)

out = __file__.replace("build_sample_deck.py", "sample-transition-sizing-deck.pptx")
prs.save(out)
print("saved", out, "| slides:", len(prs.slides))
